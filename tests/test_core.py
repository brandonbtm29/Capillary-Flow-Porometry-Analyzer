import numpy as np
import pytest

from porometry.analysis import Settings, analyze
from porometry.cft import PSI_TO_KPA, read_cft
from porometry.export import run_table, save_figure, write_summary
from porometry.qc import check
from porometry.workspace import Workspace

from _data import DATA, EXPECTED, needs_data

pytestmark = needs_data  # real lab files, kept in SyncThing (see tests/_data.py)
GOOD = DATA / "CFP-0_25MM_MIN_TEST1.CFT"  # 10-08-2026, no support mesh


def codes(name):
    a = analyze(read_cft(DATA / name))
    return {f.code for f in check(a).findings}


# ---------------------------------------------------------------- parser
def test_parse_header_and_curves():
    r = read_cft(GOOD)
    assert r.operator == "Brandon McReynolds"
    assert r.date == "10-08-2026"
    assert r.fluid == "Galwick"
    assert r.surface_tension == pytest.approx(15.9)
    assert r.diameter_cm == pytest.approx(0.635)
    assert r.thickness_cm == pytest.approx(0.075)
    assert r.tortuosity == pytest.approx(0.715)
    assert r.wet.n == 123 and r.dry.n == 132
    assert np.all(np.diff(r.wet.pressure) >= 0) and np.all(np.diff(r.dry.pressure) >= 0)
    assert r.parameters[0].startswith("B,")


# ---------------------------------------------------------------- matches CapWin's log.txt
@pytest.mark.parametrize("name,kpa,um", [
    (x["file"], x["kpa"], x["um"]) for x in EXPECTED.get("capwin_log_bubble_points", [])
])
def test_bubble_point_matches_capwin_log(name, kpa, um):
    a = analyze(read_cft(DATA / name))
    assert a.bubble_dp * PSI_TO_KPA == pytest.approx(kpa, abs=0.001)
    assert a.bubble_d == pytest.approx(um, abs=0.001)


def test_capwin_log_fixture_has_those_values():
    log = (DATA / "capwin_log.txt").read_text()
    for x in EXPECTED["capwin_log_bubble_points"]:
        assert f"{x['kpa']:.3f}" in log and f"{x['um']:.3f}" in log


def test_good_run_results():
    a = analyze(read_cft(GOOD))
    exp = EXPECTED["good_run"]
    assert a.mean_flow_dp * PSI_TO_KPA == pytest.approx(exp["mean_flow_kpa"], abs=0.2)
    assert a.mean_flow_d == pytest.approx(exp["mean_flow_um"], abs=0.002)
    assert 0.2 < a.smallest_d < 0.3
    # distribution accounts for (nearly) all the flow
    _, _, dcff = a.distribution()
    assert dcff.sum() == pytest.approx(100, abs=1)
    edges, flow_pct, _, _ = a.binned_distribution(15)
    assert flow_pct.sum() == pytest.approx(dcff.sum(), abs=1e-6)


def test_shape_factor_one_scales_diameters():
    a = analyze(read_cft(GOOD))
    b = analyze(read_cft(GOOD), Settings(shape_factor=1.0))
    assert b.bubble_d / a.bubble_d == pytest.approx(1 / 0.715)


def test_mean_flow_pore_equals_d_at_50():
    a = analyze(read_cft(GOOD))
    # CapRep finds the mean flow pore where the wet curve meets half the dry curve; the CFF = 50 %
    # crossing interpolates the ratio instead, so the two agree closely but not exactly.
    assert a.d_at_cff(50) == pytest.approx(a.mean_flow_d, rel=1e-3)


# ---------------------------------------------------------------- QC
def test_good_run_only_warns_about_wet_above_dry():
    assert codes("CFP-0_25MM_MIN_TEST1.CFT") == {"wet_above_dry"}


@pytest.mark.parametrize("name", [
    "CFP-0_25MM_MIN_TEST2.CFT", "CFP-0_25MM_MIN_TEST3.CFT", "CFP-0_70MM_MIN_TEST3.CFT",
    "CFP-10_00MM_MIN_S1_TEST1.CFT", "CFP-2_00MM_MIN_S1_TEST1.CFT",
])
def test_mesh_runs_flagged_as_leak(name):
    c = codes(name)
    assert "leak" in c and "no_mfp" in c


def test_incomplete_dry_curve():
    c = codes("CFP-0_70MM_MIN_TEST4.CFT")
    assert {"dry_missing", "below_ambient"} <= c


def test_single_point_dry_curve():
    assert "dry_missing" in codes("CFP-0_70MM_MIN_TEST1.CFT")


# ---------------------------------------------------------------- workspace / export


def test_workspace_and_exports(tmp_path):
    ws = Workspace()
    ws.add_folder(DATA)
    assert len(ws.entries) == 10 and not ws.errors
    ws.add_folder(DATA)  # duplicates ignored
    assert len(ws.entries) == 10
    s = ws.summary_frame()
    assert set(s["QC"]) <= {"pass", "info", "warn", "fail"}
    g = ws.group_frame()
    row = g[g["Group"] == "0.25 mm/min"].iloc[0]  # Condition from the file name
    assert row["Runs"] == 3 and row["Used"] == 1
    assert row["Mean flow pore mean (µm)"] == pytest.approx(EXPECTED["good_run"]["mean_flow_um"], abs=0.002)

    x = write_summary(ws, tmp_path / "out.xlsx")
    import openpyxl
    wb = openpyxl.load_workbook(x)
    assert {"Summary", "Groups", "Statistics", "QC"} <= set(wb.sheetnames) and len(wb.sheetnames) == 14
    write_summary(ws, tmp_path / "out.csv")
    assert (tmp_path / "out.csv").read_text(encoding="utf-8-sig").startswith("Run,")
    for kind in ("flow", "cff", "psd", "all"):
        for ext in ("png", "svg", "pdf"):
            p = save_figure(ws, kind, tmp_path / f"{kind}.{ext}")
            assert p.stat().st_size > 1000


def test_run_table_columns():
    t = run_table(analyze(read_cft(GOOD)))
    assert len(t) > 100 and t["CFF (%)"].iloc[-1] > 95


def test_colors_stable_after_removal():
    ws = Workspace()
    ws.add_files([DATA / "CFP-0_25MM_MIN_TEST1.CFT", DATA / "CFP-0_25MM_MIN_TEST2.CFT"])
    second = ws.entries[1].color
    ws.remove(ws.entries[0])
    assert ws.entries[0].color == second
    ws.add_files([DATA / "CFP-0_25MM_MIN_TEST3.CFT"])
    assert ws.entries[1].color not in (second,)


def test_manual_duplicate_groups():
    ws = Workspace()
    ws.add_folder(DATA)
    a = next(e for e in ws.entries if e.name == "CFP-0_25MM_MIN_TEST1")
    b = next(e for e in ws.entries if e.name == "lastdata2_2026-10-08")
    ws.set_group([a, b], "Oct-8 baseline")
    ws.reanalyze()  # overrides survive re-analysis
    assert a.group == b.group == "Oct-8 baseline"
    ws.separate([a])
    assert a.group == a.name
    ws.reset_group([a])
    assert a.group == "0.25 mm/min"


def test_startup_blip_near_zero_pressure_is_ignored():
    """CFP-PRACTICE-RUN2 logs a little flow at almost zero pressure before the real curve starts."""
    a = analyze(read_cft(DATA / "CFP-PRACTICE-RUN2.CFT"))
    assert a.bubble_dp * PSI_TO_KPA > 1          # not the 0.008 kPa blip
    exp = EXPECTED["practice_run2"]
    assert a.mean_flow_dp * PSI_TO_KPA == pytest.approx(exp["mean_flow_kpa"], abs=0.5)
    assert a.mean_flow_d == pytest.approx(exp["mean_flow_um"], abs=0.02)
