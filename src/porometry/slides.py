"""PowerPoint export: a title slide, the comparison table, then one slide per figure.

    from porometry.slides import write_deck
    write_deck("report.pptx", "Capillary flow porometry", "10-11-2026 runs",
               tables=[("Comparison", df)], figures=[("Flow vs pressure", fig_or_png_path, "caption")])

Figures are inserted as 200 dpi pictures (exactly what the notebook or app drew). A figure's own
title is left off the picture because the slide title already says it.
"""

from __future__ import annotations

import io
import math
from pathlib import Path
from typing import Iterable, Sequence, Union

import numpy as np
import pandas as pd
from matplotlib.figure import Figure
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

FigureLike = Union[Figure, str, Path]

INK = RGBColor(0x1F, 0x2A, 0x33)  # charcoal: title slide background, headings
INK2 = RGBColor(0x55, 0x5E, 0x66)  # captions
TINT = RGBColor(0xEE, 0xF1, 0xF3)  # alternate table rows
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0xEB, 0x68, 0x34)  # orange, one accent only (title slide subtitle)
FONT_HEAD = "Cambria"
FONT_BODY = "Calibri"

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6)
TITLE_TOP, TITLE_H = Inches(0.4), Inches(0.9)
BODY_TOP = Inches(1.45)
CAPTION_H = Inches(0.5)
TABLE_ROWS_PER_SLIDE = 10


def _text(slide, x, y, w, h, text, size, *, bold=False, color=INK, font=FONT_BODY, align=PP_ALIGN.LEFT,
          anchor=MSO_ANCHOR.TOP, italic=False):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    f = run.font
    f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, font
    f.color.rgb = color
    return box


def _title(slide, text):
    _text(slide, MARGIN, TITLE_TOP, W - 2 * MARGIN, TITLE_H, text, 32, bold=True, font=FONT_HEAD,
          anchor=MSO_ANCHOR.MIDDLE)


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _figure_png(fig: FigureLike) -> tuple[io.BytesIO, float]:
    """PNG bytes and width/height ratio. Titles are hidden while the picture is made, then restored."""
    if not isinstance(fig, Figure):
        data = Path(fig).read_bytes()
        from PIL import Image

        with Image.open(io.BytesIO(data)) as im:
            return io.BytesIO(data), im.width / im.height
    saved = []
    for ax in fig.axes:
        for loc in ("left", "center", "right"):
            t = ax.get_title(loc=loc)
            if t:
                saved.append((ax, loc, t))
                ax.set_title("", loc=loc)
    sup = fig._suptitle.get_text() if getattr(fig, "_suptitle", None) is not None else None
    if sup:
        fig._suptitle.set_text("")
    buf = io.BytesIO()
    try:
        fig.savefig(buf, format="png", dpi=200, bbox_inches="tight", facecolor="white")
    finally:
        for ax, loc, t in saved:
            ax.set_title(t, loc=loc)
        if sup:
            fig._suptitle.set_text(sup)
    buf.seek(0)
    from PIL import Image

    with Image.open(buf) as im:
        ratio = im.width / im.height
    buf.seek(0)
    return buf, ratio


def _column_formatter(col: pd.Series):
    """One format per column, so the decimals line up: about 3 significant figures for a typical value."""
    nums = pd.to_numeric(col, errors="coerce")
    finite = nums[np.isfinite(nums)] if nums.notna().any() else nums.iloc[:0]
    is_bool = col.map(lambda v: isinstance(v, (bool, np.bool_))).all()
    if is_bool or finite.empty or nums.notna().sum() < col.notna().sum():
        return lambda v: "-" if v is None or (isinstance(v, float) and math.isnan(v)) else (
            ("yes" if v else "no") if isinstance(v, (bool, np.bool_)) else str(v))
    if np.allclose(finite, np.round(finite, 2)):  # short exact values such as rates: 0.25, 0.7, 2, 10
        fmt = "{:g}"
    else:
        typical = float(np.median(np.abs(finite[finite != 0]))) if (finite != 0).any() else 1.0
        dec = int(min(3, max(0, 2 - math.floor(math.log10(typical)))))
        fmt = "{:,.%df}" % dec
    return lambda v: "-" if v is None or (isinstance(v, float) and math.isnan(v)) else fmt.format(float(v))


def _table_slides(prs, title: str, df: pd.DataFrame, note: str = "") -> None:
    if "Run" in df.columns:  # run names first, they get the wide column
        df = df[["Run"] + [c for c in df.columns if c != "Run"]]
    fmts = [_column_formatter(df[c]) for c in df.columns]
    pages = max(1, math.ceil(len(df) / TABLE_ROWS_PER_SLIDE))
    for page in range(pages):
        part = df.iloc[page * TABLE_ROWS_PER_SLIDE:(page + 1) * TABLE_ROWS_PER_SLIDE]
        s = _blank(prs)
        _title(s, title if pages == 1 else f"{title} ({page + 1} of {pages})")
        nrows, ncols = len(part) + 1, len(part.columns)
        width = W - 2 * MARGIN
        few = len(part) <= 6
        row_h = Inches(0.62) if few else Inches(0.48)
        body_pt = 16 if few else 14
        shape = s.shapes.add_table(nrows, ncols, MARGIN, BODY_TOP, width, Emu(row_h * nrows))
        tbl = shape.table
        # first column (run names) gets more room
        weights = [1.0] * ncols
        weights[0] = 2.2 if ncols > 1 else 1.0
        total = sum(weights)
        for j, wgt in enumerate(weights):
            tbl.columns[j].width = Emu(int(width * wgt / total))
        for i in range(nrows):
            tbl.rows[i].height = row_h
            for j in range(ncols):
                cell = tbl.cell(i, j)
                cell.margin_left = cell.margin_right = Inches(0.08)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT
                run = p.add_run()
                if i == 0:
                    run.text = str(part.columns[j])
                    run.font.bold, run.font.size = True, Pt(13)
                    run.font.color.rgb = WHITE
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = INK
                    p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
                else:
                    run.text = fmts[j](part.iloc[i - 1, j])
                    run.font.size = Pt(body_pt)
                    run.font.color.rgb = INK
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = TINT if i % 2 == 0 else WHITE
                run.font.name = FONT_BODY
        if note:
            _text(s, MARGIN, H - Inches(0.9), W - 2 * MARGIN, CAPTION_H, note, 12, color=INK2, italic=True)


def _figure_slide(prs, title: str, fig: FigureLike, caption: str = "") -> None:
    s = _blank(prs)
    _title(s, title)
    png, ratio = _figure_png(fig)
    box_w = W - 2 * MARGIN
    box_h = H - BODY_TOP - Inches(0.35) - (CAPTION_H + Inches(0.15) if caption else 0)
    w = min(box_w, int(box_h * ratio))
    h = int(w / ratio)
    x = int((W - w) / 2)
    s.shapes.add_picture(png, x, BODY_TOP, width=Emu(w), height=Emu(h))
    if caption:
        _text(s, MARGIN, BODY_TOP + h + Inches(0.15), W - 2 * MARGIN, CAPTION_H, caption, 13, color=INK2)


def write_deck(path: Union[str, Path], title: str, subtitle: str = "",
               tables: Iterable[Sequence] = (), figures: Iterable[Sequence] = (), footer: str = "") -> Path:
    """Write a 16:9 PowerPoint deck.

    tables:  (slide title, DataFrame[, note]) items; long tables continue on extra slides.
    figures: (slide title, matplotlib Figure or image path[, caption]) items, one slide each.
    """
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    s = _blank(prs)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = INK
    _text(s, MARGIN, Inches(2.4), W - 2 * MARGIN, Inches(1.5), title, 44, bold=True, color=WHITE,
          font=FONT_HEAD, anchor=MSO_ANCHOR.BOTTOM)
    if subtitle:
        _text(s, MARGIN, Inches(4.05), W - 2 * MARGIN, Inches(0.8), subtitle, 22, color=ACCENT)
    if footer:
        _text(s, MARGIN, H - Inches(1.0), W - 2 * MARGIN, Inches(0.5), footer, 13, color=RGBColor(0xC9, 0xCF, 0xD4))
    for item in tables:
        _table_slides(prs, item[0], item[1], item[2] if len(item) > 2 else "")
    for item in figures:
        _figure_slide(prs, item[0], item[1], item[2] if len(item) > 2 else "")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return path
